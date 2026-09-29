#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成一页推广 PPT（16:9）。

用法：python docs/promo/build-onepager.py
依赖：python-pptx（DSH 自带运行时里有：dependencies/python/Lib/site-packages）

设计原则（写死在这里，不靠手拖）：
- **一页只讲一件事**：插件能帮你把输入框变顺手；不铺功能清单之外的内容。
- 左文右图：左边是六条能力（用户能看懂的说法），右边留给实机截图。
- 截图可选：`docs/promo/screenshot.png` 存在就用它，不存在就画一个占位框
  （所以这个脚本在截图到位之前也能跑，生成一份能看的一页）。
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "dsh-composer-ux-onepager.pptx")
SHOT = os.path.join(HERE, "screenshot.png")

INK = RGBColor(0x1F, 0x24, 0x2E)          # 正文深色
MUTED = RGBColor(0x6B, 0x72, 0x80)        # 次要文字
ACCENT = RGBColor(0x4C, 0x6E, 0xF5)       # 主色（DSH 蓝）
ACCENT_SOFT = RGBColor(0xE8, 0xEE, 0xFD)  # 主色浅底
PANEL = RGBColor(0xF6, 0xF7, 0xF9)        # 浅灰底
CODE_BG = RGBColor(0x1B, 0x21, 0x2C)      # 代码条底色
CODE_FG = RGBColor(0xE6, 0xEA, 0xF2)

FONT = "Microsoft YaHei"
MONO = "Consolas"

FEATURES = [
    ("键位随你定", "发送键 / 换行键给预设，也能自己按一遍录下来"),
    ("右键菜单三档", "官方不介入 / 浏览器菜单 / 补上 7 项的自定义菜单"),
    ("快捷指令 + 提示词优化", "常用提示词一键插入；草稿先让模型改，逐条给出依据"),
    ("输入框下方的两个数字", "缓存命中给三位小数；本会话金额胶囊点开看明细"),
    ("终端与请求头", "Windows 默认换 Git Bash；OpenCode 请求自动带头"),
    ("一栏一个开关", "十项能力各自可关，装前可体检，设置页一栏一项"),
]


def add_text(slide, left, top, width, height, text, size=14, bold=False,
             color=INK, font=FONT, align=PP_ALIGN.LEFT, spacing=None, anchor=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    if anchor is not None:
        frame.vertical_anchor = anchor
    para = frame.paragraphs[0]
    para.alignment = align
    if spacing is not None:
        para.line_spacing = spacing
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def add_rect(slide, left, top, width, height, fill, line=None, radius=None):
    from pptx.enum.shapes import MSO_SHAPE

    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    if radius is not None:
        shape.adjustments[0] = radius
    return shape


def main():
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])  # 空白版式

    # 顶部主色条 + 标题
    add_rect(slide, 0, 0, 13.333, 0.18, ACCENT)
    add_text(slide, 0.7, 0.5, 7.4, 0.9, "dsh-composer-ux", size=40, bold=True, color=INK)
    add_text(slide, 0.72, 1.42, 7.4, 0.5,
             "DeepSeek Harness 网页版 · 输入体验增强插件", size=17, color=ACCENT)
    add_text(slide, 0.72, 1.92, 7.3, 0.5,
             "让输入框顺手起来：键位、右键菜单、快捷指令、提示词优化、金额与统计。", size=13, color=MUTED)

    # 六条能力（左栏两列 × 三行）
    top0, row_h = 2.62, 1.02
    for index, (title, desc) in enumerate(FEATURES):
        col = index % 2
        row = index // 2
        left = 0.72 + col * 3.72
        top = top0 + row * row_h
        add_rect(slide, left, top + 0.08, 0.09, 0.62, ACCENT, radius=0.45)
        add_text(slide, left + 0.22, top, 3.35, 0.32, title, size=13.5, bold=True, color=INK)
        add_text(slide, left + 0.22, top + 0.28, 3.4, 0.62, desc, size=10.5, color=MUTED, spacing=1.15)

    # 底部安装条
    add_rect(slide, 0.72, 5.92, 7.3, 0.66, CODE_BG, radius=0.18)
    add_text(slide, 0.98, 6.02, 6.9, 0.5,
             "dsh plugin --profile <profile> add dsh-composer-ux",
             size=13, color=CODE_FG, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.72, 6.72, 7.5, 0.4,
             "装完重启 DSH 即可；设置页多出「输入体验」一栏，十项能力各自可关。",
             size=11, color=MUTED)

    # 右侧截图区
    shot_left, shot_top, shot_w, shot_h = 8.42, 0.72, 4.2, 6.36
    if os.path.exists(SHOT):
        slide.shapes.add_picture(SHOT, Inches(shot_left), Inches(shot_top), width=Inches(shot_w))
        add_text(slide, shot_left, 7.02, shot_w, 0.3, "实机截图", size=10, color=MUTED, align=PP_ALIGN.CENTER)
    else:
        add_rect(slide, shot_left, shot_top, shot_w, shot_h, PANEL, radius=0.04)
        add_text(slide, shot_left + 0.3, shot_top + 2.6, shot_w - 0.6, 1.2,
                 "（截图占位）\n把实机截图另存为 docs/promo/screenshot.png\n再跑一次本脚本即可替换",
                 size=11, color=MUTED, align=PP_ALIGN.CENTER, spacing=1.3)

    # 页脚
    add_text(slide, 0.72, 7.06, 7.4, 0.3,
             "0.12.0 · MIT · github.com/fangwen9527/dsh-composer-ux", size=10, color=MUTED)

    deck.save(OUT)
    print("已生成", OUT)
    print("截图：", "已包含" if os.path.exists(SHOT) else "占位（docs/promo/screenshot.png 还不存在）")


if __name__ == "__main__":
    main()
